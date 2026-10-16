"""Import first-year timetable workbook into Turso. Usage: python import_timetable.py workbook.xlsx"""
import hashlib, json, os, sys
from datetime import date, datetime, time
from pathlib import Path
from libsql_client import create_client_sync
from openpyxl import load_workbook

def env(name):
    value=os.getenv(name,"").strip()
    if not value: raise SystemExit("Missing environment variable: "+name)
    return value
def val(x):
    if x is None: return None
    if isinstance(x,datetime): return x.date().isoformat()
    if isinstance(x,date): return x.isoformat()
    if isinstance(x,time): return x.strftime("%H:%M")
    return x if isinstance(x,(str,int,float)) else str(x)
def record(headers, row): return {str(h).strip():v for h,v in zip(headers,row) if h}
def key(prefix,payload):
    raw=json.dumps({k:val(v) for k,v in payload.items()},ensure_ascii=False,sort_keys=True,separators=(",",":"))
    return prefix+":"+hashlib.sha256(raw.encode()).hexdigest()
url,token=env("TURSO_DATABASE_URL"),env("TURSO_AUTH_TOKEN")
if url.startswith("libsql://"): url="https://"+url[len("libsql://"):]
if not url.startswith("https://"): raise SystemExit("TURSO_DATABASE_URL must start libsql:// or https://")
if len(sys.argv)!=2: raise SystemExit('Usage: python import_timetable.py "workbook.xlsx"')
path=Path(sys.argv[1]).expanduser()
if not path.is_file(): raise SystemExit("Workbook not found: "+str(path))
wb=load_workbook(path,read_only=True,data_only=True)
required={"Timetable","Translation Rules","Exam Timetable","AI Interpretation Rules"}
if not required.issubset(wb.sheetnames): raise SystemExit("Workbook missing sheets: "+", ".join(required-set(wb.sheetnames)))
db=create_client_sync(url=url,auth_token=token)
schema=[
"""CREATE TABLE IF NOT EXISTS timetable_events(source_key TEXT PRIMARY KEY,event_date TEXT NOT NULL,day TEXT NOT NULL,start_time TEXT NOT NULL,end_time TEXT NOT NULL,subject TEXT NOT NULL,activity_type TEXT NOT NULL,groups_text TEXT,room_text TEXT,group_room_association TEXT,verification_status TEXT,review_note TEXT,source_page TEXT,source_cell TEXT,source_text TEXT,year_level TEXT NOT NULL DEFAULT 'First year',updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""",
"CREATE INDEX IF NOT EXISTS idx_timetable_date_time ON timetable_events(event_date,start_time)",
"CREATE INDEX IF NOT EXISTS idx_timetable_subject ON timetable_events(subject,activity_type)",
"""CREATE TABLE IF NOT EXISTS exam_events(source_key TEXT PRIMARY KEY,event_date TEXT NOT NULL,day TEXT NOT NULL,assessment_type TEXT NOT NULL,subject TEXT NOT NULL,activity_component TEXT,year_level TEXT NOT NULL DEFAULT 'First year',source_text TEXT,review_note TEXT,updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""",
"CREATE INDEX IF NOT EXISTS idx_exam_date ON exam_events(event_date)",
"""CREATE TABLE IF NOT EXISTS translation_aliases(id INTEGER PRIMARY KEY AUTOINCREMENT,entity_type TEXT NOT NULL,canonical_value TEXT NOT NULL,alias TEXT NOT NULL,language_script TEXT,matching_rule TEXT,notes TEXT,UNIQUE(entity_type,canonical_value,alias))""",
"CREATE INDEX IF NOT EXISTS idx_alias_type ON translation_aliases(entity_type)",
"CREATE TABLE IF NOT EXISTS interpretation_rules(rule_area TEXT NOT NULL,rule TEXT NOT NULL,PRIMARY KEY(rule_area,rule))"]
for sql in schema: db.execute(sql)
counts={}
ws=wb["Timetable"]; it=ws.iter_rows(values_only=True); headers=next(it); n=0
for row in it:
 r=record(headers,row)
 if not r.get("Date") or not r.get("Subject"): continue
 fields=("Date","Day","Start Time","End Time","Subject","Activity Type","Groups","Room","Group-Room Association","Verification Status","Review Note","Source Page","Source Cell","Source Text")
 p={f:r.get(f) for f in fields}; k=key("timetable",p)
 args=(k,val(p["Date"]),str(p["Day"] or ""),val(p["Start Time"]) or "",val(p["End Time"]) or "",str(p["Subject"]),str(p["Activity Type"] or ""),val(p["Groups"]),val(p["Room"]),val(p["Group-Room Association"]),val(p["Verification Status"]),val(p["Review Note"]),val(p["Source Page"]),val(p["Source Cell"]),val(p["Source Text"]))
 db.execute("""INSERT INTO timetable_events(source_key,event_date,day,start_time,end_time,subject,activity_type,groups_text,room_text,group_room_association,verification_status,review_note,source_page,source_cell,source_text) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(source_key) DO UPDATE SET event_date=excluded.event_date,day=excluded.day,start_time=excluded.start_time,end_time=excluded.end_time,subject=excluded.subject,activity_type=excluded.activity_type,groups_text=excluded.groups_text,room_text=excluded.room_text,group_room_association=excluded.group_room_association,verification_status=excluded.verification_status,review_note=excluded.review_note,source_page=excluded.source_page,source_cell=excluded.source_cell,source_text=excluded.source_text,updated_at=CURRENT_TIMESTAMP""",args); n+=1
counts["timetable_events"]=n
ws=wb["Exam Timetable"]; it=ws.iter_rows(values_only=True); headers=next(it); n=0
for row in it:
 r=record(headers,row)
 if not r.get("Date") or not r.get("Subject"): continue
 fields=("Date","Day","Assessment Type","Subject","Activity / Exam Component","Year","Source Text","Review Note"); p={f:r.get(f) for f in fields}; k=key("exam",p)
 db.execute("""INSERT INTO exam_events(source_key,event_date,day,assessment_type,subject,activity_component,year_level,source_text,review_note) VALUES(?,?,?,?,?,?,?,?,?) ON CONFLICT(source_key) DO UPDATE SET event_date=excluded.event_date,day=excluded.day,assessment_type=excluded.assessment_type,subject=excluded.subject,activity_component=excluded.activity_component,year_level=excluded.year_level,source_text=excluded.source_text,review_note=excluded.review_note,updated_at=CURRENT_TIMESTAMP""",(k,val(p["Date"]),str(p["Day"] or ""),str(p["Assessment Type"] or "Not specified"),str(p["Subject"]),val(p["Activity / Exam Component"]),str(p["Year"] or "First year"),val(p["Source Text"]),val(p["Review Note"]))); n+=1
counts["exam_events"]=n
ws=wb["Translation Rules"]; it=ws.iter_rows(values_only=True); headers=next(it); n=0
for row in it:
 r=record(headers,row); alias=r.get("Alias / Variant"); canon=r.get("Canonical Value"); typ=r.get("Entity Type")
 if not alias or not canon or not typ: continue
 db.execute("""INSERT INTO translation_aliases(entity_type,canonical_value,alias,language_script,matching_rule,notes) VALUES(?,?,?,?,?,?) ON CONFLICT(entity_type,canonical_value,alias) DO UPDATE SET language_script=excluded.language_script,matching_rule=excluded.matching_rule,notes=excluded.notes""",(str(typ),str(canon),str(alias),val(r.get("Language / Script")),val(r.get("Matching Rule")),val(r.get("Notes")))); n+=1
counts["translation_aliases"]=n
ws=wb["AI Interpretation Rules"]; it=ws.iter_rows(values_only=True); headers=next(it); n=0
for row in it:
 r=record(headers,row); area,rule=r.get("Rule Area"),r.get("Rule")
 if not area or not rule: continue
 db.execute("INSERT INTO interpretation_rules(rule_area,rule) VALUES(?,?) ON CONFLICT(rule_area,rule) DO NOTHING",(str(area),str(rule))); n+=1
counts["interpretation_rules"]=n
for table,n in counts.items(): print(f"{table}: processed {n}; database total {db.execute('SELECT COUNT(*) FROM '+table).rows[0][0]}")
wb.close(); db.close()
print("Import complete. Re-running the same workbook is safe.")
